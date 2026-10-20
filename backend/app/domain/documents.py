"""Lebenslauf upload and text extraction. Shared by the HTTP API and the bot.

* The file type is detected from content (magic bytes), never from the name/extension.
* DOCX is a ZIP: the total uncompressed size is checked before parsing (zip bombs).
* A PDF without a text layer (scan) is stored with status `no_text`; the user is asked
  for a DOCX or a text PDF. No OCR (see ARCHITECTURE §8a).
* One current CV per user: a new upload replaces the previous one (data minimisation).

Functions here flush but never commit — the caller owns the transaction.
"""

import asyncio
import base64
import io
import re
import uuid
import zipfile
from dataclasses import dataclass
from typing import Any

import pdfplumber
import structlog
from docx import Document as DocxDocument
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Document, DocumentStatus, User
from app.domain.accounts import AccountError, audit

log = structlog.get_logger()

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 10
MIN_TEXT_CHARS = 200  # less than this from a PDF means "probably a scan"

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class Extracted:
    content_type: str
    text: str | None


def detect_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist():
                    if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED_BYTES:
                        raise AccountError("file_too_large")
                    return DOCX
        except zipfile.BadZipFile:
            pass
    raise AccountError("unsupported_file_type")


def _normalize(text: str) -> str:
    text = text.replace("\u00ad", "")  # soft hyphens from PDF exports
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages[:MAX_PDF_PAGES]
        return "\n\n".join(page.extract_text() or "" for page in pages)


def _docx_text(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Lebenslauf templates often put everything into tables.
    for table in doc.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if cell.text and (not cells or cells[-1] != cell.text):  # merged cells repeat
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract(data: bytes) -> Extracted:
    """CPU-bound: call through `asyncio.to_thread`."""
    content_type = detect_type(data)
    try:
        raw = _pdf_text(data) if content_type == PDF else _docx_text(data)
    except AccountError:
        raise
    except Exception as exc:  # corrupt or encrypted file
        log.warning("document_parse_failed", content_type=content_type, error=type(exc).__name__)
        raise AccountError("file_unreadable") from exc
    text = _normalize(raw)
    if content_type == PDF and len(text) < MIN_TEXT_CHARS:
        return Extracted(content_type, None)
    return Extracted(content_type, text or None)


async def upload_cv(db: AsyncSession, user: User, *, filename: str, data: bytes) -> Document:
    if not data:
        raise AccountError("file_empty")
    if len(data) > MAX_FILE_BYTES:
        raise AccountError("file_too_large")
    extracted = await asyncio.to_thread(extract, data)

    await db.execute(delete(Document).where(Document.user_id == user.id, Document.kind == "cv"))
    document = Document(
        user_id=user.id,
        kind="cv",
        filename=filename[:255] or "lebenslauf",
        content_type=extracted.content_type,
        size_bytes=len(data),
        file=data,
        text=extracted.text,
        status=DocumentStatus.EXTRACTED if extracted.text else DocumentStatus.NO_TEXT,
    )
    db.add(document)
    await audit(db, user.id, "document.cv_uploaded")
    await db.flush()
    log.info(
        "cv_uploaded",
        user_id=str(user.id),
        content_type=extracted.content_type,
        size_bytes=len(data),
        status=document.status,
        text_chars=len(extracted.text or ""),
    )
    return document


async def current_cv(db: AsyncSession, user_id: uuid.UUID) -> Document | None:
    return await db.scalar(
        select(Document)
        .where(Document.user_id == user_id, Document.kind == "cv")
        .order_by(Document.created_at.desc())
        .limit(1)
    )


async def delete_cv(db: AsyncSession, user: User) -> None:
    await db.execute(delete(Document).where(Document.user_id == user.id, Document.kind == "cv"))
    await audit(db, user.id, "document.cv_deleted")
    await db.flush()


async def export_documents(db: AsyncSession, user_id: uuid.UUID) -> list[dict[str, Any]]:
    """GDPR export: metadata, extracted text and the original file (base64)."""
    documents = (await db.scalars(select(Document).where(Document.user_id == user_id))).all()
    return [
        {
            "kind": d.kind,
            "filename": d.filename,
            "content_type": d.content_type,
            "size_bytes": d.size_bytes,
            "status": d.status,
            "uploaded_at": d.created_at.isoformat(),
            "text": d.text,
            "file_base64": base64.b64encode(d.file).decode(),
        }
        for d in documents
    ]
