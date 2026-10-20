import base64
import io
import zipfile

import pytest
from docx import Document as DocxDocument
from httpx import AsyncClient
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Document
from app.domain import documents
from tests.conftest import PASSWORD, register

CV_LINES = [
    "Max Mustermann",
    "Berufserfahrung: Koch, Restaurant am Hafen, Hamburg, 2019-2024",
    "Ausbildung: Koch, Berufsschule Hamburg, 2016-2019",
    "Sprachen: Deutsch C1, Englisch B2",
    "Kenntnisse: A-la-carte-Kueche, HACCP, Warenwirtschaft",
]


def make_pdf(lines: list[str]) -> bytes:
    """Minimal valid one-page PDF with a Helvetica text layer (ASCII only)."""
    content = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({x}) Tj T*" for x in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out


def make_docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Lebenslauf — Anna Schmidt")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "2018–2024"
    table.cell(0, 1).text = "Pflegefachfrau, Klinikum Köln"
    table.cell(1, 0).text = "Sprachen"
    table.cell(1, 1).text = "Deutsch C2, Türkisch Muttersprache"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_text_is_extracted() -> None:
    result = documents.extract(make_pdf(CV_LINES))
    assert result.content_type == documents.PDF
    assert result.text is not None
    assert "Restaurant am Hafen" in result.text and "HACCP" in result.text


def test_pdf_without_text_layer_is_a_scan() -> None:
    assert documents.extract(make_pdf([])).text is None


def test_docx_paragraphs_and_tables() -> None:
    result = documents.extract(make_docx())
    assert result.content_type == documents.DOCX
    assert result.text is not None
    assert "Anna Schmidt" in result.text
    assert "2018–2024 | Pflegefachfrau, Klinikum Köln" in result.text
    assert "Türkisch" in result.text


@pytest.mark.parametrize(
    "data",
    [
        b"\x89PNG\r\n\x1a\n" + b"0" * 100,  # image
        b"plain text pretending to be a CV",
        b"PK\x03\x04 not really a zip",
    ],
)
def test_unsupported_types(data: bytes) -> None:
    with pytest.raises(documents.AccountError, match="unsupported_file_type"):
        documents.extract(data)


def test_zip_without_word_document_is_rejected() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(documents.AccountError, match="unsupported_file_type"):
        documents.extract(buf.getvalue())


def test_zip_bomb_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents, "MAX_UNZIPPED_BYTES", 10_000)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", "0" * 50_000)
    with pytest.raises(documents.AccountError, match="file_too_large"):
        documents.extract(buf.getvalue())


def test_corrupt_pdf() -> None:
    with pytest.raises(documents.AccountError, match="file_unreadable"):
        documents.extract(b"%PDF-1.4 garbage")


async def upload(
    client: AsyncClient, data: bytes, name: str = "lebenslauf.pdf"
) -> dict[str, object]:
    resp = await client.post(
        "/documents/cv", files={"file": (name, data, "application/octet-stream")}
    )
    assert resp.status_code == 201, resp.text
    result: dict[str, object] = resp.json()
    return result


async def test_upload_replace_and_delete(client: AsyncClient, db: AsyncSession) -> None:
    await register(client)
    assert (await client.get("/documents/cv")).status_code == 404

    first = await upload(client, make_pdf(CV_LINES))
    assert first["status"] == "extracted"
    assert first["content_type"] == documents.PDF
    assert int(str(first["text_chars"])) > 100

    second = await upload(client, make_docx(), "cv.docx")
    assert (await client.get("/documents/cv")).json()["id"] == second["id"]
    assert await db.scalar(select(func.count()).select_from(Document)) == 1  # replaced

    assert (await client.delete("/documents/cv")).status_code == 204
    assert (await client.get("/documents/cv")).status_code == 404


async def test_scan_is_stored_with_no_text(client: AsyncClient) -> None:
    await register(client)
    assert (await upload(client, make_pdf([])))["status"] == "no_text"


async def test_upload_errors(client: AsyncClient, monkeypatch: pytest.MonkeyPatch) -> None:
    await register(client)
    resp = await client.post(
        "/documents/cv", files={"file": ("x.png", b"\x89PNG....", "image/png")}
    )
    assert (resp.status_code, resp.json()["detail"]) == (415, "unsupported_file_type")

    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 100)
    resp = await client.post("/documents/cv", files={"file": ("cv.pdf", make_pdf(CV_LINES))})
    assert (resp.status_code, resp.json()["detail"]) == (413, "file_too_large")


async def test_upload_requires_login(client: AsyncClient) -> None:
    resp = await client.post("/documents/cv", files={"file": ("cv.pdf", make_pdf(CV_LINES))})
    assert resp.status_code == 401


async def test_file_and_text_encrypted_at_rest(client: AsyncClient, db: AsyncSession) -> None:
    await register(client)
    await upload(client, make_pdf(CV_LINES), "Max_Mustermann_CV.pdf")
    row = (await db.execute(text("SELECT filename_enc, file_enc, text_enc FROM documents"))).one()
    for value in row:
        assert value.startswith("v1:")
        assert "Mustermann" not in value
    assert b"Mustermann" not in base64.b64decode(row.file_enc.split(":", 1)[1] + "==")


async def test_export_and_account_deletion(client: AsyncClient, db: AsyncSession) -> None:
    await register(client)
    pdf = make_pdf(CV_LINES)
    await upload(client, pdf)

    [exported] = (await client.get("/me/export")).json()["documents"]
    assert exported["filename"] == "lebenslauf.pdf"
    assert "HACCP" in exported["text"]
    assert base64.b64decode(exported["file_base64"]) == pdf

    await client.post("/me/delete", json={"password": PASSWORD})
    assert await db.scalar(select(func.count()).select_from(Document)) == 0
